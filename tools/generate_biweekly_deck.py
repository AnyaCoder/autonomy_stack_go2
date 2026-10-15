from pathlib import Path
import subprocess

from PIL import Image

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path(__file__).resolve().parents[1] / "docs" / "Go2双周进展汇报.pptx"
ASSETS = Path(__file__).resolve().parents[1] / "docs" / "assets"
W, H = 13.333, 7.5
BG = "F5F7F6"
INK = "17262B"
MUTED = "5B6B70"
TEAL = "087E75"
MINT = "DCEFEA"
BLUE = "346B91"
PALE_BLUE = "E4EEF4"
AMBER = "D58B35"
PALE_AMBER = "F8ECD9"
RED = "B85C4A"
WHITE = "FFFFFF"
LINE = "D8E0DD"
FONT = "Microsoft YaHei"


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def rect(slide, x, y, w, h, color, line=None, radius=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.color.rgb = rgb(line or color)
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def text(slide, value, x, y, w, h, size=16, color=INK, bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.02):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = value
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def rich_lines(slide, lines, x, y, w, h, size=14, color=INK, gap=8):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.level = 0
        if isinstance(line, tuple):
            lead, body, lead_color = line
            r = p.add_run()
            r.text = lead
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.bold = True
            r.font.color.rgb = rgb(lead_color)
            r = p.add_run()
            r.text = body
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.color.rgb = rgb(color)
        else:
            p.text = line
            p.font.name = FONT
            p.font.size = Pt(size)
            p.font.color.rgb = rgb(color)
    return box


def new_slide(prs, title, kicker, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(BG)
    rect(slide, 0, 0, W, 0.09, TEAL)
    text(slide, kicker.upper(), 0.62, 0.38, 6.5, 0.25, 10, TEAL, True)
    text(slide, title, 0.62, 0.78, 12, 0.55, 26, INK, True)
    rect(slide, 0.62, 1.48, 12.08, 0.012, LINE)
    text(slide, f"GO2  /  双周进展                                      {number:02d}",
         0.62, 7.13, 12.1, 0.2, 9, MUTED)
    return slide


def label(slide, value, x, y, w, fill, color):
    rect(slide, x, y, w, 0.29, fill, fill, True)
    text(slide, value, x + 0.08, y + 0.035, w - 0.16, 0.19, 10, color, True,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)


def bullet(slide, value, x, y, w, size=14, color=INK):
    rect(slide, x, y + 0.09, 0.07, 0.07, TEAL, TEAL, True)
    text(slide, value, x + 0.19, y, w - 0.19, 0.47, size, color)


def crop_pdf_figure(pdf_path, page, prefix, box):
    ASSETS.mkdir(parents=True, exist_ok=True)
    output = ASSETS / f"{prefix}.jpg"
    if output.exists():
        return output
    temp_prefix = Path("/tmp") / f"{prefix}-figure-page"
    subprocess.run(["pdftoppm", "-f", str(page), "-singlefile", "-r", "110",
                    "-png", str(pdf_path), str(temp_prefix)], check=True)
    source = Image.open(f"{temp_prefix}.png").convert("RGB")
    source.crop(box).save(output, quality=90, optimize=True)
    return output


def add_picture(slide, path, x, y, w, h):
    image = Image.open(path)
    iw, ih = image.size
    temp_image = ASSETS / f"_crop_{Path(path).stem}.jpg"
    if image.mode not in ("RGB", "L"):
        image = image.convert("RGB")
    image.save(temp_image, quality=92, optimize=True)
    scale = min(w / iw, h / ih)
    draw_w, draw_h = iw * scale, ih * scale
    draw_x, draw_y = x + (w - draw_w) / 2, y + (h - draw_h) / 2
    slide.shapes.add_picture(str(temp_image), Inches(draw_x), Inches(draw_y),
                             Inches(draw_w), Inches(draw_h))
    border = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(draw_x), Inches(draw_y),
                                    Inches(draw_w), Inches(draw_h))
    border.fill.background()
    border.line.color.rgb = rgb(LINE)


prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
prs.core_properties.title = "Go2 双周进展汇报"
prs.core_properties.subject = "论文调研与实机部署"
prs.core_properties.author = ""

# 1. Cover
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = rgb(INK)
rect(slide, 0.72, 0.85, 0.1, 4.9, TEAL)
text(slide, "双周进展汇报", 1.12, 1.18, 7.8, 0.7, 35, WHITE, True)
text(slide, "论文调研  ×  Go2 实机部署", 1.15, 2.12, 7.8, 0.45, 21, "B9DAD2", True)
text(slide, "LM-Nav  /  RT-2  /  OmniNav", 1.15, 3.0, 6.6, 0.32, 14, "D3DDDA")
text(slide, "完成 L1 IMU 校准与 Point-LIO SLAM 实机运行", 1.15, 3.64,
     7.4, 0.36, 15, WHITE)
text(slide, "2026.10", 1.15, 5.52, 2.4, 0.3, 13, "AFC0BB", True)

# Abstract sensor and navigation motif
for x, y, label_text, color in [
    (9.2, 1.32, "语言指令", TEAL),
    (10.9, 2.22, "语义目标", BLUE),
    (9.45, 3.16, "定位建图", AMBER),
    (11.15, 4.18, "闭环执行", TEAL),
]:
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.34), Inches(0.34))
    circle.fill.solid()
    circle.fill.fore_color.rgb = rgb(color)
    circle.line.color.rgb = rgb(color)
    text(slide, label_text, x + 0.48, y + 0.035, 1.6, 0.25, 12, WHITE, True)
for x1, y1, x2, y2 in [(9.37, 1.66, 11.07, 2.22), (11.07, 2.56, 9.62, 3.16),
                         (9.62, 3.5, 11.32, 4.18)]:
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = rgb("55716F")
    conn.line.width = Pt(2)
text(slide, "感知  →  建图  →  导航", 9.2, 5.45, 3.2, 0.35, 13, "B9DAD2", True)
rect(slide, 0, 7.27, W, 0.23, TEAL)

# 2. Snapshot
slide = new_slide(prs, "本阶段进展", "阶段总览", 2)
col_x = [0.68, 4.86, 9.04]
blocks = [
    ("论文调研", "3 篇", TEAL, MINT,
     ["梳理语言理解、视觉语义与机器人动作的连接方式", "归纳长时序导航中的记忆、搜索与执行闭环"]),
    ("实机部署", "Go2 + L1", BLUE, PALE_BLUE,
     ["克隆并编译 autonomy_stack_go2", "完成实机通信检查与 L1 IMU 校准"]),
    ("运行验证", "SLAM", AMBER, PALE_AMBER,
     ["Point-LIO 接入 L1 点云与 IMU 数据", "RViz 可查看点云、里程计与建图过程"]),
]
for x, (heading, metric, accent, pale, items) in zip(col_x, blocks):
    rect(slide, x, 1.9, 3.62, 4.62, WHITE, LINE, True)
    rect(slide, x, 1.9, 3.62, 0.08, accent)
    text(slide, heading, x + 0.25, 2.25, 3.1, 0.35, 18, INK, True)
    text(slide, metric, x + 0.25, 2.88, 3.05, 0.5, 25, accent, True)
    rect(slide, x + 0.25, 3.57, 3.1, 0.015, LINE)
    bullet(slide, items[0], x + 0.25, 3.92, 3.1, 13)
    bullet(slide, items[1], x + 0.25, 4.8, 3.1, 13)
text(slide, "当前阶段已建立实机感知与定位建图基础，导航闭环进入后续验证。",
     0.75, 6.68, 11.8, 0.28, 14, MUTED, True, align=PP_ALIGN.CENTER)

# 3. Papers
slide = new_slide(prs, "三篇论文：从指令理解走向长时序执行", "论文调研", 3)
pdf_dir = Path("/home/anyacoder/Documents/xwechat_files/wxid_r801ixy5z1zn22_7498/msg/file/2026-10")
figures = [
    (0.68, "LM-Nav", "语言解析地标，视觉匹配目标，导航策略执行。", TEAL,
     crop_pdf_figure(pdf_dir / "LM-Nav Robotic Navigation with Large Pre-Trained Models of Language, Vision, and Action.pdf",
                     2, "lmnav-framework-v2", (160, 108, 748, 230))),
    (4.86, "RT-2", "把动作表示为 token，与视觉语言数据共同训练。", BLUE,
     crop_pdf_figure(pdf_dir / "RT-2 Vision-Language-Action Models Transfer Web Knowledge to Robotic Control.pdf",
                     2, "rt2-framework-v2", (92, 125, 820, 390))),
    (9.04, "OmniNav", "动态场景记忆、证据更新与可操作位姿联合规划。", AMBER,
     crop_pdf_figure(pdf_dir / "OmniNav- Robust Long-Horizon Target Navigation in Dynamic Environments.pdf",
                     1, "omninav-framework-v2", (68, 250, 875, 615))),
]
for x, name, summary, accent, figure_path in figures:
    text(slide, name, x + 0.05, 1.9, 3.56, 0.34, 18, accent, True)
    add_picture(slide, figure_path, x + 0.05, 2.42, 3.5, 2.52)
    text(slide, summary, x + 0.05, 5.15, 3.52, 0.76, 12, INK, True)
    text(slide, "来源：论文框架图", x + 0.05, 6.23, 3.52, 0.24, 10, MUTED)
text(slide, "共同脉络：语义理解 → 空间决策 → 机器人执行与反馈",
     0.76, 6.64, 11.5, 0.28, 14, TEAL, True, align=PP_ALIGN.CENTER)

# 4. Synthesis
slide = new_slide(prs, "研究结论：构建可验证的导航闭环", "对项目的启发", 4)
steps = [
    (0.74, "任务表达", "自然语言 / 目标点", TEAL, MINT),
    (3.88, "语义落地", "地标与场景目标", BLUE, PALE_BLUE),
    (7.02, "空间规划", "地图、可达性、路径", AMBER, PALE_AMBER),
    (10.16, "运动执行", "局部控制与反馈", TEAL, MINT),
]
for x, title, desc, color, pale in steps:
    rect(slide, x, 2.05, 2.4, 1.28, pale, pale, True)
    text(slide, title, x + 0.14, 2.29, 2.12, 0.31, 17, color, True,
         align=PP_ALIGN.CENTER)
    text(slide, desc, x + 0.1, 2.76, 2.2, 0.28, 11, MUTED,
         align=PP_ALIGN.CENTER)
for x in [3.24, 6.38, 9.52]:
    text(slide, "→", x, 2.45, 0.48, 0.35, 22, MUTED, True, align=PP_ALIGN.CENTER)
rect(slide, 0.74, 3.86, 11.82, 0.015, LINE)
text(slide, "三篇工作共同指向", 0.82, 4.18, 2.5, 0.3, 13, MUTED, True)
rich_lines(slide, [
    ("分层决策：", "语言模型负责目标表达，导航模块负责空间与运动约束。", TEAL),
    ("持续更新：", "定位、地图和目标状态需要随观测与执行结果迭代。", BLUE),
    ("逐级验证：", "先验证传感器与定位，再验证路径跟随，最后扩展语义任务。", AMBER),
], 0.82, 4.62, 11.2, 1.6, 15, INK, 12)

# 5. Go2 stack and achieved results
slide = new_slide(prs, "Go2 实机：已跑通感知与建图链路", "系统部署", 5)
nodes = [
    (0.7, 1.9, 2.22, 1.0, "Unitree Go2 + L1", "实机传感器", PALE_BLUE, BLUE),
    (3.54, 1.9, 2.25, 1.0, "DDS 数据链路", "IMU / 点云样本", MINT, TEAL),
    (6.42, 1.9, 2.25, 1.0, "传感器变换", "外参配置与坐标系", PALE_AMBER, AMBER),
    (9.3, 1.9, 2.25, 1.0, "Point-LIO", "里程计与地图", MINT, TEAL),
]
for x, y, w, h, title, desc, pale, color in nodes:
    rect(slide, x, y, w, h, pale, pale, True)
    text(slide, title, x + 0.1, y + 0.16, w - 0.2, 0.28, 14, color, True,
         align=PP_ALIGN.CENTER)
    text(slide, desc, x + 0.1, y + 0.58, w - 0.2, 0.22, 10, MUTED,
         align=PP_ALIGN.CENTER)
for x in [2.96, 5.83, 8.71]:
    text(slide, "→", x, 2.21, 0.44, 0.35, 20, MUTED, True, align=PP_ALIGN.CENTER)
rviz_path = Path(__file__).resolve().parents[1] / "img" / "rviz_full.jpg"
add_picture(slide, rviz_path, 6.92, 3.37, 5.36, 2.55)
text(slide, "仓库 RViz 界面示例（仿真截图）", 6.92, 6.0, 5.36, 0.23, 10, MUTED)
rect(slide, 0.7, 3.14, 11.62, 0.015, LINE)
text(slide, "实机验证", 0.76, 3.48, 1.5, 0.3, 14, TEAL, True)
results = [
    (0.76, "通信检查", "L1 IMU 与点云话题样本正常"),
    (4.15, "IMU 校准", "完成标定并生成参数文件"),
]
for x, heading, desc in results:
    text(slide, "✓", x, 4.05, 0.3, 0.35, 19, TEAL, True)
    text(slide, heading, x + 0.4, 4.04, 2.7, 0.29, 15, INK, True)
    text(slide, desc, x + 0.4, 4.48, 3.0, 0.48, 11, MUTED)
text(slide, "Point-LIO SLAM", 0.76, 5.36, 2.5, 0.29, 15, INK, True)
text(slide, "实机运行：点云、里程计与轨迹可视化", 0.76, 5.78, 5.7, 0.3, 12, MUTED)
text(slide, "下一步：验证外参与时间同步，再接入路径跟随。",
     0.78, 6.42, 5.7, 0.42, 13, RED, True)

# 6. Next steps
slide = new_slide(prs, "下一阶段：从建图走向稳定导航", "后续两周安排", 6)
timeline = [
    (0.76, "01", "稳定性复测", "重复采集同一路线，检查地图连续性、漂移与系统启动稳定性。", TEAL, MINT),
    (4.64, "02", "参数验证", "按录包回放评估 IMU-LiDAR 外参与时间同步，再更新实机配置。", BLUE, PALE_BLUE),
    (8.52, "03", "导航闭环", "接入局部规划与目标点跟随，在安全场地完成低速往返测试。", AMBER, PALE_AMBER),
]
for x, number, title, desc, color, pale in timeline:
    rect(slide, x, 2.05, 3.25, 2.4, WHITE, LINE, True)
    text(slide, number, x + 0.22, 2.27, 0.65, 0.42, 22, color, True)
    rect(slide, x + 0.22, 2.88, 2.8, 0.035, pale)
    text(slide, title, x + 0.22, 3.15, 2.8, 0.32, 17, INK, True)
    text(slide, desc, x + 0.22, 3.63, 2.8, 0.56, 12, MUTED)
text(slide, "验收顺序", 0.8, 5.02, 1.5, 0.3, 13, MUTED, True)
rect(slide, 0.8, 5.52, 11.25, 0.055, LINE)
for x, name, color in [(0.8, "重复建图", TEAL), (4.63, "定位稳定", BLUE),
                       (8.52, "目标点到达", AMBER)]:
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(5.37), Inches(0.34), Inches(0.34))
    dot.fill.solid()
    dot.fill.fore_color.rgb = rgb(color)
    dot.line.color.rgb = rgb(color)
    text(slide, name, x + 0.42, 5.84, 2.55, 0.28, 12, INK, True)
text(slide, "阶段目标：形成可复现的 Go2 感知—定位—路径跟随验证流程。",
     0.8, 6.32, 11.7, 0.34, 15, TEAL, True)

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(OUT)
